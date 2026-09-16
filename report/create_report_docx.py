from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "report" / "P0_NL_to_SQL_Bakeoff_Report_Draft.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0].cells
    for index, text in enumerate(headers):
        header[index].width = Inches(widths[index])
        header[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(header[index], "203864")
        set_cell_border(header[index])
        set_cell_margins(header[index])
        paragraph = header[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8.5)
    for row_number, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cells[index])
            set_cell_margins(cells[index])
            if row_number % 2:
                set_cell_shading(cells[index], "EEF3F8")
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(text))
            run.font.size = Pt(8.3)
    return table


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.04
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        run.bold = True
    paragraph.add_run(text)
    return paragraph


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor(32, 32, 32)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

heading = styles["Heading 1"]
heading.font.name = "Aptos Display"
heading._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
heading.font.size = Pt(12.5)
heading.font.bold = True
heading.font.color.rgb = RGBColor(0, 0, 0)
heading.paragraph_format.space_before = Pt(7)
heading.paragraph_format.space_after = Pt(3)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
title.add_run("Project 0 Natural Language to SQL Bake Off")
remove_paragraph_border(title)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(8)
run = subtitle.add_run("CS496 AI Engineering  |  Report draft  |  16 September 2026")
run.italic = True
run.font.size = Pt(9)

identity = doc.add_paragraph()
identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
identity.paragraph_format.space_after = Pt(8)
identity.add_run("Student: ____________________    Team: ____________________").font.size = Pt(9)

add_heading(doc, "Objective")
add_body(
    doc,
    "The benchmark measures how accurately and efficiently language models convert English bookstore questions into SQLite queries. The approved task is evaluated with automatic execution-based scoring so that different SQL formulations can receive credit when they return equivalent results. The completed local-model result is 31 correct answers out of 54 test items, or 57.4 percent. Results for the two required API models remain pending course API access.",
)

add_heading(doc, "Dataset and database")
add_body(
    doc,
    "A deterministic synthetic online bookstore was selected because it supports both basic SQL and multi-table reasoning. Its seven tables cover customers, categories, products, orders, order items, payments, and returns. The database contains 120 customers, 80 products, 300 orders, 750 order items, 360 payments, and 60 returns. Seed 496 makes regeneration reproducible. Designed edge cases include customers with no orders, products never ordered, cancelled orders, failed payments, and partial returns.",
)
add_body(
    doc,
    "The benchmark contains 12 development items and 54 test items. Development items were used to verify the pipeline and clarify the prompt before measurement. The fixed test set contains 18 easy, 20 medium, and 16 hard questions, exceeding the required minimum of 50 test items. Each item stores an English question and reference SQL, but only the English question is sent to the model.",
)

add_heading(doc, "Method")
add_body(
    doc,
    "For every item, the runner combines the shared schema prompt with one English question. The model returns one read-only SQL query. The scorer executes the generated SQL and the hidden reference SQL against the same SQLite database, then compares their result columns and rows. It supports ordered and unordered answers, duplicates, NULL values, numeric tolerance, empty results, timeouts, and execution errors. Invalid SQL and non-equivalent results count as incorrect.",
)
add_body(
    doc,
    "The shared prompt defines the schema, relationships, status values, money units, and return quantities. The same fixed prompt, item order, temperature, output limit, and parser will be used for all three models. Calls are sequential so latency measurements reflect one request at a time.",
)

add_heading(doc, "Models and settings")
add_table(
    doc,
    ["Role", "Exact model", "Temperature", "Output limit", "Status"],
    [
        ["Local open weights", "qwen2.5-coder:3b", "0", "256 tokens", "Complete"],
        ["Top API", "Pending course access", "0 planned", "256 planned", "Pending"],
        ["Cheap API", "Pending course access", "0 planned", "256 planned", "Pending"],
    ],
    [1.55, 2.2, 1.0, 1.15, 1.05],
)

doc.add_page_break()

add_heading(doc, "Local model results")
add_body(
    doc,
    "The local model ran through Ollama on a Windows laptop with 24 GB DDR4 memory and an NVIDIA GeForce RTX 3050 Ti Laptop GPU with 4 GB dedicated memory. Ollama used a 4,096-token context window, seed 496, streaming disabled, and one sequential request per question. The development run reached 10 out of 12 after prompt clarification. The fixed test run then produced the measurements below.",
)

add_table(
    doc,
    ["Metric", "Measured value", "Meaning"],
    [
        ["Accuracy", "31 of 54  (57.4%)", "Equivalent query results"],
        ["p50 latency", "3.39 seconds", "Median complete response time"],
        ["p95 latency", "4.67 seconds", "95 percent of calls at or below this value"],
        ["Output speed", "65.9 tokens per second", "Generated tokens divided by generation time"],
        ["Hosted API cost", "None", "Hardware and electricity not measured"],
    ],
    [1.45, 1.65, 3.85],
)

add_heading(doc, "Interpretation")
add_body(
    doc,
    "The local model solved a majority of the fixed test questions while running entirely on the available laptop. Its 57.4 percent accuracy establishes the baseline for the later API comparison. Incorrect items included executable queries that returned different results and invalid queries that referenced nonexistent columns. These failures are counted rather than repaired because the goal is to measure the model's first response under fixed conditions.",
)
add_body(
    doc,
    "Latency was stable enough for a small sequential benchmark: the median response took 3.39 seconds and the p95 response took 4.67 seconds. The reported 65.9 tokens per second measures generation speed rather than full end-to-end response time. No hosted API fee was incurred, although local hardware and electricity costs were not estimated.",
)

add_heading(doc, "Comparison status")
add_body(
    doc,
    "A final quality, speed, and cost recommendation cannot yet be made because the top and cheap API models have not been selected or executed. When course access is provided, each API model will receive the same 54 test items and fixed prompt. Per-item correctness, latency, and token use will be appended to results/per_item.csv. Model-level accuracy, p50 and p95 latency, total cost, and estimated cost per 1,000 questions will then be calculated in results/summary.csv.",
)

add_heading(doc, "Limitations and next steps")
add_body(
    doc,
    "Execution equivalence on one synthetic database can accept queries that happen to produce the same result on that data, so question wording and reference SQL still require careful review. Exact GPU offload was not recorded. The remaining work is to obtain the two API models, complete identical runs, calculate actual API costs from measured token use and current prices, finish the three-model comparison, and update this draft without changing the fixed test set or local results.",
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Project 0 report draft  |  API comparison pending")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "Project 0 Natural Language to SQL Bake Off"
doc.core_properties.subject = "Draft two-page benchmark report"
doc.core_properties.author = "CS496 Project Team"
doc.save(OUTPUT)
print(OUTPUT)
